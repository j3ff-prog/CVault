"""
services/cv_extractor.py
Extracts plain text from uploaded CV files (PDF, DOCX, TXT).
"""
import os
import tempfile


def extract_text(file_storage) -> str:
    filename = file_storage.filename.lower()
    suffix = os.path.splitext(filename)[1]

    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        file_storage.save(tmp.name)
        tmp_path = tmp.name

    try:
        if filename.endswith(".pdf"):
            return _extract_pdf(tmp_path)
        elif filename.endswith(".docx"):
            return _extract_docx(tmp_path)
        elif filename.endswith(".doc"):
            return _extract_doc(tmp_path)
        elif filename.endswith(".txt"):
            return _extract_txt(tmp_path)
        else:
            raise ValueError(
                f"Unsupported file type '{suffix}'. "
                "Please upload a PDF, Word (.docx), or TXT file."
            )
    finally:
        os.unlink(tmp_path)


def _extract_pdf(path: str) -> str:
    try:
        import fitz
        doc = fitz.open(path)
        pages = [page.get_text("text").strip() for page in doc if page.get_text("text").strip()]
        doc.close()
        result = "\n\n".join(pages)
        if not result.strip():
            raise ValueError(
                "Could not extract text from this PDF — it may be scanned or image-only. "
                "Please paste your CV as text instead."
            )
        return result
    except ImportError:
        raise RuntimeError("PyMuPDF not installed. Run: pip install pymupdf")


def _extract_docx(path: str) -> str:
    try:
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        result = "\n".join(parts)
        if not result.strip():
            raise ValueError("The Word document appears empty. Please check the file or paste your CV as text.")
        return result
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")


def _extract_doc(path: str) -> str:
    import subprocess
    try:
        result = subprocess.run(["antiword", path], capture_output=True, text=True, timeout=10)
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout.strip()
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass
    raise ValueError(
        "Old .doc format is not fully supported. "
        "Please save your CV as .docx or paste the text directly."
    )


def _extract_txt(path: str) -> str:
    for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            with open(path, "r", encoding=enc) as f:
                text = f.read()
            if text.strip():
                return text.strip()
        except (UnicodeDecodeError, IOError):
            continue
    raise ValueError("Could not read the text file. Please check the file encoding.")
